import os
import sqlite3
from datetime import datetime
from typing import List, Dict, Any, Optional

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None


def get_main_excel_path(db_path: str) -> str:
    base_dir = os.path.dirname(os.path.abspath(db_path))
    base_name = os.path.splitext(os.path.basename(db_path))[0]
    return os.path.join(base_dir, f"{base_name}.xlsx")


def ensure_main_excel(db_path: str) -> str:
    path = get_main_excel_path(db_path)
    if Workbook is None:
        return path
    if not os.path.exists(path):
        export_main_excel(db_path, path)
    return path


def initialize_db(db_path: Optional[str] = None) -> str:
    if db_path is None:
        db_path = "inventory_billing.db"

    conn = sqlite3.connect(db_path)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            price REAL NOT NULL,
            barcode TEXT,
            stock_quantity INTEGER NOT NULL DEFAULT 0
        )
        """
    )

    columns = [row[1] for row in conn.execute("PRAGMA table_info(products)")]
    if "barcode" not in columns:
        conn.execute("ALTER TABLE products ADD COLUMN barcode TEXT")
    if "stock_quantity" not in columns:
        conn.execute("ALTER TABLE products ADD COLUMN stock_quantity INTEGER NOT NULL DEFAULT 0")

    conn.execute(
        "UPDATE products SET stock_quantity = 0 WHERE stock_quantity IS NULL"
    )
    conn.commit()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS sales (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_name TEXT,
            customer_phone TEXT,
            sale_date TEXT NOT NULL,
            subtotal REAL NOT NULL,
            discount_rate REAL NOT NULL DEFAULT 0,
            discount_amount REAL NOT NULL DEFAULT 0,
            vat_rate REAL NOT NULL DEFAULT 0,
            vat_amount REAL NOT NULL DEFAULT 0,
            total_amount REAL NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS sale_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sale_id INTEGER NOT NULL,
            product_id INTEGER NOT NULL,
            quantity INTEGER NOT NULL,
            unit_price REAL NOT NULL,
            line_total REAL NOT NULL,
            FOREIGN KEY (sale_id) REFERENCES sales(id),
            FOREIGN KEY (product_id) REFERENCES products(id)
        )
        """
    )
    conn.execute(
        "INSERT OR IGNORE INTO users (username, password) VALUES (?, ?)",
        ("admin", "admin"),
    )
    conn.commit()
    conn.close()
    return db_path


def add_product(
    db_path: str,
    name: str,
    price: float,
    barcode: Optional[str] = None,
    stock_quantity: int = 0,
) -> int:
    conn = sqlite3.connect(db_path)
    cursor = conn.execute(
        "INSERT INTO products (name, price, barcode, stock_quantity) VALUES (?, ?, ?, ?)",
        (name.strip(), float(price), (barcode or "").strip() or None, int(stock_quantity)),
    )
    conn.commit()
    product_id = cursor.lastrowid
    conn.close()
    if Workbook is not None:
        export_main_excel(db_path, get_main_excel_path(db_path))
    return product_id


def get_all_products(db_path: str) -> List[Dict[str, Any]]:
    conn = sqlite3.connect(db_path)
    rows = conn.execute(
        "SELECT id, name, price, barcode, stock_quantity FROM products ORDER BY name"
    ).fetchall()
    conn.close()
    return [
        {
            "id": row[0],
            "name": row[1],
            "price": row[2],
            "barcode": row[3],
            "stock_quantity": row[4],
        }
        for row in rows
    ]


def update_product(
    db_path: str,
    product_id: int,
    name: str,
    price: float,
    barcode: Optional[str] = None,
    stock_quantity: Optional[int] = None,
) -> bool:
    conn = sqlite3.connect(db_path)
    if stock_quantity is None:
        cursor = conn.execute(
            "UPDATE products SET name = ?, price = ?, barcode = ? WHERE id = ?",
            (name.strip(), float(price), (barcode or "").strip() or None, product_id),
        )
    else:
        cursor = conn.execute(
            "UPDATE products SET name = ?, price = ?, barcode = ?, stock_quantity = ? WHERE id = ?",
            (name.strip(), float(price), (barcode or "").strip() or None, int(stock_quantity), product_id),
        )
    conn.commit()
    conn.close()
    if Workbook is not None:
        export_main_excel(db_path, get_main_excel_path(db_path))
    return cursor.rowcount > 0


def delete_product(db_path: str, product_id: int) -> bool:
    conn = sqlite3.connect(db_path)
    cursor = conn.execute("DELETE FROM products WHERE id = ?", (product_id,))
    conn.commit()
    conn.close()
    if Workbook is not None:
        export_main_excel(db_path, get_main_excel_path(db_path))
    return cursor.rowcount > 0


def search_products(db_path: str, search_term: str) -> List[Dict[str, Any]]:
    term = f"%{search_term.strip().lower()}%"
    conn = sqlite3.connect(db_path)
    rows = conn.execute(
        "SELECT id, name, price, barcode, stock_quantity FROM products WHERE lower(name) LIKE ? ORDER BY name",
        (term,),
    ).fetchall()
    conn.close()
    return [
        {
            "id": row[0],
            "name": row[1],
            "price": row[2],
            "barcode": row[3],
            "stock_quantity": row[4],
        }
        for row in rows
    ]


def get_product_by_barcode(db_path: str, barcode: str) -> Optional[Dict[str, Any]]:
    conn = sqlite3.connect(db_path)
    row = conn.execute(
        "SELECT id, name, price, barcode, stock_quantity FROM products WHERE lower(barcode) = ?",
        (barcode.strip().lower(),),
    ).fetchone()
    conn.close()
    if row is None:
        return None
    return {
        "id": row[0],
        "name": row[1],
        "price": row[2],
        "barcode": row[3],
        "stock_quantity": row[4],
    }


def authenticate_user(db_path: str, username: str, password: str) -> bool:
    conn = sqlite3.connect(db_path)
    row = conn.execute(
        "SELECT 1 FROM users WHERE username = ? AND password = ?",
        (username.strip(), password),
    ).fetchone()
    conn.close()
    return row is not None


def calculate_item_total(price: float, quantity: int) -> float:
    return float(price) * int(quantity)


def calculate_grand_total(items: List[Dict[str, Any]]) -> float:
    return sum(calculate_item_total(item["price"], item["quantity"]) for item in items)


def record_sale(
    db_path: str,
    items: List[Dict[str, Any]],
    customer_name: str = "",
    customer_phone: str = "",
    discount_rate: float = 0,
    vat_rate: float = 0,
) -> int:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    subtotal = 0.0
    sale_items: List[tuple] = []

    for item in items:
        product_id = int(item["product_id"])
        quantity = int(item["quantity"])
        product = conn.execute(
            "SELECT id, name, price, stock_quantity FROM products WHERE id = ?",
            (product_id,),
        ).fetchone()
        if product is None:
            conn.close()
            raise ValueError("Product not found")
        if quantity > int(product["stock_quantity"]):
            conn.close()
            raise ValueError(f"Not enough stock for {product['name']}")

        unit_price = float(item.get("price", product["price"]))
        line_total = unit_price * quantity
        subtotal += line_total
        sale_items.append((product_id, quantity, unit_price, line_total))

    discount_amount = subtotal * (float(discount_rate) / 100.0)
    taxable_amount = subtotal - discount_amount
    vat_amount = taxable_amount * (float(vat_rate) / 100.0)
    total_amount = taxable_amount + vat_amount

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor = conn.execute(
        """
        INSERT INTO sales (
            customer_name, customer_phone, sale_date, subtotal, discount_rate,
            discount_amount, vat_rate, vat_amount, total_amount
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            customer_name.strip() or "Walk-in",
            customer_phone.strip() or "",
            now,
            subtotal,
            float(discount_rate),
            discount_amount,
            float(vat_rate),
            vat_amount,
            total_amount,
        ),
    )
    sale_id = cursor.lastrowid

    for product_id, quantity, unit_price, line_total in sale_items:
        conn.execute(
            "INSERT INTO sale_items (sale_id, product_id, quantity, unit_price, line_total) VALUES (?, ?, ?, ?, ?)",
            (sale_id, product_id, quantity, unit_price, line_total),
        )
        conn.execute(
            "UPDATE products SET stock_quantity = stock_quantity - ? WHERE id = ?",
            (quantity, product_id),
        )

    conn.commit()
    conn.close()
    if Workbook is not None:
        export_main_excel(db_path, get_main_excel_path(db_path))
    return sale_id


def get_sales_report(db_path: str) -> List[Dict[str, Any]]:
    conn = sqlite3.connect(db_path)
    rows = conn.execute(
        "SELECT id, customer_name, customer_phone, sale_date, total_amount FROM sales ORDER BY sale_date DESC"
    ).fetchall()
    conn.close()
    return [
        {
            "id": row[0],
            "customer_name": row[1],
            "customer_phone": row[2],
            "sale_date": row[3],
            "total_amount": row[4],
        }
        for row in rows
    ]


def get_daily_income_report(db_path: str) -> Dict[str, Any]:
    conn = sqlite3.connect(db_path)
    row = conn.execute(
        """
        SELECT COUNT(*) AS sales_count, COALESCE(SUM(total_amount), 0) AS total_sales
        FROM sales
        WHERE date(sale_date) = date('now')
        """
    ).fetchone()
    conn.close()
    return {
        "date": datetime.now().strftime("%Y-%m-%d"),
        "sales_count": row[0],
        "total_sales": row[1],
    }


def export_products_to_excel(db_path: str, file_path: str) -> None:
    if Workbook is None:
        raise ImportError("openpyxl is required for Excel export")

    products = get_all_products(db_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["ID", "Name", "Price", "Barcode", "Stock Quantity"])
    for product in products:
        sheet.append([
            product["id"],
            product["name"],
            product["price"],
            product.get("barcode", ""),
            product.get("stock_quantity", 0),
        ])
    workbook.save(file_path)


def export_sales_to_excel(db_path: str, file_path: str) -> None:
    if Workbook is None:
        raise ImportError("openpyxl is required for Excel export")

    conn = sqlite3.connect(db_path)
    rows = conn.execute(
        "SELECT id, customer_name, customer_phone, sale_date, subtotal, discount_rate, discount_amount, vat_rate, vat_amount, total_amount FROM sales ORDER BY sale_date DESC"
    ).fetchall()
    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append([
        "ID",
        "Customer Name",
        "Customer Phone",
        "Sale Date",
        "Subtotal",
        "Discount %",
        "Discount Amount",
        "VAT %",
        "VAT Amount",
        "Total Amount",
    ])
    for row in rows:
        sheet.append(row)
    workbook.save(file_path)


def export_main_excel(db_path: str, file_path: str) -> None:
    if Workbook is None:
        raise ImportError("openpyxl is required for Excel export")

    workbook = Workbook()

    # Products worksheet
    products = get_all_products(db_path)
    products_sheet = workbook.active
    products_sheet.title = "Products"
    products_sheet.append(["ID", "Name", "Price", "Barcode", "Stock Quantity"])
    for product in products:
        products_sheet.append([
            product["id"],
            product["name"],
            product["price"],
            product.get("barcode", ""),
            product.get("stock_quantity", 0),
        ])

    # Sales worksheet
    sales_sheet = workbook.create_sheet(title="Sales")
    conn = sqlite3.connect(db_path)
    rows = conn.execute(
        "SELECT id, customer_name, customer_phone, sale_date, subtotal, discount_rate, discount_amount, vat_rate, vat_amount, total_amount FROM sales ORDER BY sale_date DESC"
    ).fetchall()
    conn.close()
    sales_sheet.append([
        "ID",
        "Customer Name",
        "Customer Phone",
        "Sale Date",
        "Subtotal",
        "Discount %",
        "Discount Amount",
        "VAT %",
        "VAT Amount",
        "Total Amount",
    ])
    for row in rows:
        sales_sheet.append(row)

    workbook.save(file_path)
